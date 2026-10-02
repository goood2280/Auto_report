"""Offline service contract tests. No mail, S3, or corporate APIs."""
import io
import json
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from pptx import Presentation

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import Main as main
import My_Function as mf
import Scheduler as scheduler
from anomaly_engine import analyze_commonality

NOW=pd.Timestamp('2026-09-19 09:30')

def formatter():
    return pd.DataFrame([dict(ALIAS=item,ITEMID=item,CAT2=cat,CATEGORY='REAL',UNIT='V',
        SPECLOW=0.,SPECHIGH=10.,**{'REPORT DIRECTION':'BOTH','REPORT ORDER':i+1,'REPORT LOG SCALE':False})
        for i,(item,cat) in enumerate([('CURRENT','Electrical'),('STABLE','Electrical'),('LEAKAGE','Leakage')])])

def frame():
    rows=[]
    for lot in range(36):
        stamp=NOW-pd.Timedelta(days=36-lot) if lot<24 else NOW-pd.Timedelta(hours=36-lot)
        for wafer in range(1,3):
            for shot in range(7):
                rows.append(dict(fab_lot_id=f'L{lot:03}.1',root_lot_id=f'L{lot:03}',wafer_id=wafer,
                    step_id='S1',step_seq='P1',temperature=25,tkout_time=stamp,
                    chip_x_pos=shot-3,chip_y_pos=wafer-1,flat_zone=0,
                    CURRENT=5+shot*.03+(7 if lot>=24 else np.sin(lot)*.01),
                    STABLE=5+shot*.03,LEAKAGE=5+shot*.03+(2 if lot>=24 and shot==6 else 0)))
    return pd.DataFrame(rows)

@pytest.fixture(autouse=True)
def offline(tmp_path,monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv('AUTO_REPORT_OPS_ROOT',str(tmp_path/'ops'))
    monkeypatch.setenv('AUTO_REPORT_TEMP_DIR',str(tmp_path/'basis'))
    monkeypatch.setattr(main.GLOBAL_CONFIG,'settings',dict(vehicle='TEST',viewing_period=45,trend_tkout_agg={},DB=str(tmp_path)))
    monkeypatch.setattr(main.GLOBAL_CONFIG,'load_from_yaml',lambda vehicle:None)
    def forbidden(*a,**k):raise AssertionError('External operation forbidden')
    monkeypatch.setattr(main.requests,'request',forbidden)
    durable_mail=main._durable_mail
    monkeypatch.setattr(main,'_durable_mail',forbidden)
    return durable_mail

def ml_file(path,product='TEST'):
    data=frame()[['root_lot_id','wafer_id']].drop_duplicates()
    data['KNOB_ETCH']='A';data.to_parquet(path/f'ML_TABLE_{product}.parquet',index=False)

def settings(tmp_path,service='daily_trend'):
    return dict(service=service,products=['TEST'],ml_table_dir=str(tmp_path),chart_dpi=120,
        report_now=NOW,highlight_since=NOW-pd.Timedelta(days=1),influence_enabled=False)

def test_daily_rolling_day_ignores_publication_checkpoint():
    data=frame()
    data.loc[data.root_lot_id.eq('L024'),'tkout_time']=NOW-pd.Timedelta(hours=23)
    cfg=dict(service='daily_trend',highlight_since=NOW-pd.Timedelta(days=1),_published_observations=set())
    entry=mf.daily_trend_entries(data,formatter(),'TEST',cfg,NOW)[0]
    assert entry['recent_lots']==12
    assert entry['points'].loc[entry['points']['_recent'],'_dc_time'].min()<NOW.normalize()
    assert not entry['points'].loc[entry['points']['_dc_time']<=NOW-pd.Timedelta(days=1),'_recent'].any()

def test_daily_uses_exact_production_detector_and_no_basis_writes(tmp_path):
    entries=mf.daily_trend_entries(frame(),formatter(),'TEST',settings(tmp_path),NOW)
    mf.daily_auto_findings(entries,formatter())
    assert not (tmp_path/'basis').exists()
    assert next(e for e in entries if e['item']=='CURRENT')['auto_findings']
    assert not next(e for e in entries if e['item']=='STABLE')['auto_findings']
    direct=analyze_commonality(frame(),'L024.1',{a:{} for a in formatter().ALIAS},formatter().set_index('ALIAS'),config=main.GLOBAL_CONFIG,persist_basis=False)
    expected={(f['item'],f['type'],f['severity']) for f in direct if f['severity'] in ('CRITICAL','WARNING') and not f.get('_excl_unless')}
    actual={(e['item'],f['type'],f['severity']) for e in entries for f in e['auto_findings'] if f['lot']=='L024.1'}
    assert actual==expected

@pytest.mark.parametrize('service',['daily_trend','mlmode'])
def test_missing_ml_skips_before_reading_or_sending(service,tmp_path,monkeypatch):
    monkeypatch.setattr(main,'daily_trend_load',lambda *a:pytest.fail('Must skip before DB load'))
    cfg=settings(tmp_path,service);cfg.pop('report_now');cfg.pop('highlight_since');cfg['recipients']=['offline@example.test']
    result=main._daily_trend_report(dict(id='missing',service=service,settings=cfg,now=NOW.timestamp(),send=True))
    assert result['status']=='skipped' and result['parts']==[]
    assert result['coverage'][0]['status']=='missing_ml_table'
    assert not mf.ops_list('trend_publication_checkpoints')

def test_daily_delivery_is_frozen_and_partial_missing_product_skips(tmp_path,monkeypatch):
    ml_file(tmp_path)
    monkeypatch.setattr(main.pd,'read_csv',lambda *a,**k:formatter())
    monkeypatch.setattr(main,'daily_trend_load',lambda *a:frame())
    cfg=dict(products=['TEST','ABSENT'],ml_table_dir=str(tmp_path),recipients=['offline@example.test'],chart_dpi=120)
    calls=[]
    monkeypatch.setattr(main,'_durable_mail',lambda *a:calls.append(a) or 'failed')
    req=dict(id='retry',service='daily_trend',settings=cfg,now=NOW.timestamp(),send=True)
    first=main._daily_trend_report(req)
    assert first['status']=='failed' and first['items']==3
    assert not mf.ops_list('trend_publication_checkpoints')
    assert any(c['status']=='missing_ml_table' for c in first['coverage'])
    monkeypatch.setattr(main,'daily_trend_load',lambda *a:pytest.fail('Must reuse frozen artifacts'))
    monkeypatch.setattr(main,'_durable_mail',lambda *a:calls.append(a) or 'sent')
    second=main._daily_trend_report(req)
    assert second['status']=='sent' and second['parts']==first['parts']
    assert calls[0][0]==calls[-1][0] and str(calls[0][4]).endswith('.pptx')
    assert mf.ops_list('trend_publication_checkpoints')

def test_catalog_edit_changes_artifacts_and_publication_checkpoint(tmp_path,monkeypatch):
    """The same daily ID must not reuse a report made for a different item selection."""
    import yaml
    ml_file(tmp_path)
    path=tmp_path/'report_items.yaml'
    monkeypatch.setattr(main.pd,'read_csv',lambda *a,**k:formatter())
    monkeypatch.setattr(main,'daily_trend_load',lambda *a:frame())
    monkeypatch.setattr(main,'_daily_trend_chart',lambda *a:b'unused mocked chart')
    monkeypatch.setattr(main,'_daily_trend_pack',lambda entries,*a:[('<html>preview</html>',b'ppt',len(entries))])
    cfg=dict(products=['TEST'],items_file=str(path),ml_table_dir=str(tmp_path),recipients=[])
    req=dict(id='same-daily-id',service='daily_trend',settings=cfg,now=NOW.timestamp(),send=False)
    path.write_text(yaml.safe_dump({'version':1,'daily_trend':{'TEST':{'CURRENT':{}}}}),encoding='utf-8')
    first=main._daily_trend_report(req)
    assert first['items']==1
    path.write_text(yaml.safe_dump({'version':1,'daily_trend':{'TEST':{'CURRENT':{},'STABLE':{}}}}),encoding='utf-8')
    second=main._daily_trend_report(req)
    assert second['items']==2
    assert second['id']!=first['id'] and second['manifest']!=first['manifest']
    assert second['checkpoint_key']!=first['checkpoint_key']
    assert second['source_fingerprint']['report_items']!=first['source_fingerprint']['report_items']
    assert not mf.ops_list('trend_publication_checkpoints')

@pytest.mark.parametrize('service',['daily_trend','mlmode'])
def test_explicit_empty_catalog_skips_before_loading_measurements(tmp_path,monkeypatch,service):
    import yaml
    ml_file(tmp_path)
    path=tmp_path/'report_items.yaml'
    path.write_text(yaml.safe_dump({'version':1,service:{'TEST':{}}}),encoding='utf-8')
    monkeypatch.setattr(main.pd,'read_csv',lambda *a,**k:formatter())
    monkeypatch.setattr(main,'daily_trend_load',lambda *a:pytest.fail('No selected items; must not load DB'))
    result=main._daily_trend_report(dict(id='empty',service=service,
        settings=dict(products=['TEST'],items_file=str(path),ml_table_dir=str(tmp_path)),now=NOW.timestamp(),send=False))
    assert result['status']=='skipped' and result['coverage'][0]['status']=='no_items'

def test_summary_links_html_and_ppt_and_size_limits(tmp_path):
    cfg=settings(tmp_path)
    entries=mf.daily_auto_findings(mf.daily_trend_entries(frame(),formatter(),'TEST',cfg,NOW),formatter())
    for e in entries:e['png']=main._daily_trend_chart(e,cfg)
    parts=main._daily_trend_pack(entries,cfg,'SYNTHETIC · Daily Trend')
    assert sum(n for _,_,n in parts)==3
    for body,ppt,count in parts:
        assert len(body.encode())<1_000_000 and len(ppt)<10_000_000
        links=re.findall('href="#(item-[^"]+)"',body)
        assert links and all(f'id="{anchor}"' in body for anchor in links)
        assert body.index('daily-findings')<body.index('trend-category')
        deck=Presentation(io.BytesIO(ppt))
        assert len(deck.slides)>=3
        assert any('hlinkClick' in shape._element.xml for shape in deck.slides[0].shapes)

def test_ml_independent_effect_threshold_and_budget(tmp_path):
    cfg=settings(tmp_path,'mlmode');cfg.update(modules=['distribution_shift','spread_change','spike_rate'],diagnostic_modules=[])
    entries=mf.daily_trend_entries(frame(),formatter(),'TEST',cfg,NOW)
    selected,analysis=mf.ml_trend_select(entries,cfg)
    assert any(e['item']=='CURRENT' for e in selected)
    assert all(e['item']!='STABLE' for e in selected)
    assert all(f['effect']>=f['minimum_effect'] for e in selected for f in e['ml_findings'])
    # 효과 임계(모듈별 minimum_effect)를 극단으로 올리면 탐지가 사라진다(네거티브 컨트롤)
    cfg.update(rank_effect_min=1e9, distribution_min_distance=2.0, rate_increase=2.0)
    entries=mf.daily_trend_entries(frame(),formatter(),'TEST',cfg,NOW)
    selected_none,_=mf.ml_trend_select(entries,cfg)
    assert not selected_none
    cfg.update(rank_effect_min=.33, distribution_min_distance=.3, rate_increase=.15)
    cfg['analysis_max_tests']=1
    entries=mf.daily_trend_entries(frame(),formatter(),'TEST',cfg,NOW)
    _,limited=mf.ml_trend_select(entries,cfg)
    assert limited['statistical_tests']==1 and limited['budget_limited']
    assert all(any('partial' in w for w in e['warnings']) for e in entries if not e['ml_test_count'])

def test_watchdog_product_status_csv_and_no_s3(tmp_path,monkeypatch):
    now=NOW.timestamp()
    for product,reason,email,status in [('SENT','발행 대기','sent','success'),('EXCLUDED','report_making=False','',''),('FAILED','발행 대기','failed','failed')]:
        key=product+'_L_S1'
        mf.ops_put('measurements',key,dict(prime_key=key,vehicle=product,lot='L',step='S1',tkout_time='t',last_seen=now-10,reason=reason))
        if status:mf.ops_put('reports',key,dict(prime_key=key,vehicle=product,lot='L',step='S1',tkout_time='t',started=now-10,status=status,generated=True,saved=True,email=email))
    calls=[];monkeypatch.setattr(main,'_durable_mail',lambda *a:calls.append(a) or 'sent')
    result=main._watchdog_report(dict(id='offline',now=now,window_start=now-86400,send=True,
        settings=dict(mail_vehicle='TEST',recipients=['offline@example.test']),health=dict(state='healthy',message='정상')))
    body=Path(result['html']).read_text(encoding='utf-8');csv=pd.read_csv(result['csv'])
    assert result['status']=='sent' and str(calls[0][4]).endswith('.csv')
    assert {'발행 대상 조건','발행 모드','다음 확인'}<=set(csv.columns)
    assert body.index('id="products"')<body.index('id="actions"')
    assert '정상 처리' in body and 'report_making=False' in body and '발행 실패' in body
    assert result['attention_count']==1

@pytest.mark.parametrize(('http_status','mail_status','exit_code'),[(200,'sent',0),(503,'unknown',1)])
def test_default_auto_entrypoint_records_delivery_outcome(tmp_path,monkeypatch,offline,http_status,mail_status,exit_code):
    """Exercise AUTO CLI, ledger and delivery; only query/render work and HTTP are replaced."""
    monkeypatch.setattr(sys,'argv',['Main.py','TEST'])
    monkeypatch.setattr(main,'_RUN',None)
    monkeypatch.setenv('AUTO_REPORT_RUN_ID','auto-offline')
    result_path=tmp_path/'result.json'
    monkeypatch.setenv('AUTO_REPORT_RESULT_PATH',str(result_path))
    monkeypatch.setenv('AUTO_REPORT_EXECUTION_WAIT_SEC','0')
    monkeypatch.setattr(main,'_drain_uploads',lambda **kwargs:None)
    monkeypatch.setattr(main,'shutdown_chart_pool',lambda:None)
    cfg=main.GLOBAL_CONFIG
    cfg.settings.update(report_making=True,DB_Setting_mode=False,specific_dc_layer=False,
                        use_email_send=True,use_s3_upload=False,email_receiver=['AUTO TEAM'])
    original_settings=dict(cfg.settings)
    html=tmp_path/'auto.html';html.write_text('<p>offline report</p>',encoding='utf-8')
    ppt=tmp_path/'auto.pptx';ppt.write_bytes(b'offline attachment')
    recipient=dict(email='offline@example.test',seq=1)
    groups=[]
    monkeypatch.setattr(main,'_trigger_receivers',lambda path,group:groups.append(group) or [recipient])
    # Restore the real delivery function disabled by the default offline fixture.
    monkeypatch.setattr(main,'_durable_mail',offline)
    sends=[]
    def transport(method,url,**kwargs):
        records=mf.ops_list('mail')
        assert len(records)==1 and records[0]['status']=='sending'
        assert records[0]['identity']=='TEST_L1_S1|AUTO|2026-09-19 09:30:00'
        assert method=='POST' and kwargs['files'][0][1][1]==b'offline attachment'
        sends.append(kwargs)
        return type('Response',(),dict(status_code=http_status,text='offline'))()
    monkeypatch.setattr(main.requests,'request',transport)
    observed=pd.DataFrame([dict(prime_key='TEST_L1_S1',lot_id='L1',dc_step_id='S1',
                                tkout_time='2026-09-19 09:30:00',dc_done=True,wafer_id=1)])
    def report_work(command):
        assert command==dict(argument='TEST',kind='legacy',recipient=None)
        assert main._apply_command_settings(command,cfg) is None
        assert cfg.settings==original_settings  # AUTO retains product delivery settings.
        main._RUN.data['vehicle']='TEST'
        main._observe_measurements(observed,observed,cfg)
        report=main._RUN.begin('TEST','L1','S1',None)
        report['email']=main._send_report_files(str(html),str(ppt),cfg.get('email_receiver'),'Offline AUTO',report['id'])
        main._RUN.finish_report('success')  # Unknown delivery must still make the run fail.
    monkeypatch.setattr(main,'_main_impl',report_work)
    assert main.main()==exit_code
    result=json.loads(result_path.read_text(encoding='utf-8'))
    assert result['argument']=='TEST' and result['vehicle']=='TEST'
    assert result['status']==('success' if exit_code==0 else 'failed')
    report=mf.ops_get('reports',result['reports'][0])
    assert report['mode']=='AUTO' and report['email']==mail_status
    assert report['status']==('success' if exit_code==0 else 'unknown')
    measurement=mf.ops_get('measurements','TEST_L1_S1')
    assert measurement['reason']=='발행 대기' and measurement['run_id']=='auto-offline'
    # Rechecking the same AUTO identity never submits sent or uncertain mail again.
    assert main._send_report_files(str(html),str(ppt),cfg.get('email_receiver'),'Offline AUTO',report['id'])==mail_status
    assert len(sends)==1 and groups==['AUTO TEAM','AUTO TEAM']
