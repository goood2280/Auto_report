"""ML mode 인자 스크리닝·리포트 구조와 메일 이미지 한도(Attach file count is over 10 방지) 계약 테스트."""
import io
import re
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import Main as main
import My_Function as mf

NOW = pd.Timestamp('2026-09-19 09:30')


def wafer_frame(effect, seed=1):
    """10 lot × 12 wafer × 5 shot. effect 별로 한 인자만 반응값에 영향을 준다."""
    rng = np.random.default_rng(seed)
    rows = []
    for lot in range(10):
        for wafer in range(1, 13):
            cd, ox = rng.uniform(0, 10), rng.uniform(18, 26)
            knob = 'A' if wafer % 2 else 'B'
            base = 100 + rng.normal(0, 2)
            if effect == 'r2':
                base += 3 * cd
            if effect == 'tail' and ox < 21 and rng.random() < .6:
                base -= 25
            if effect == 'knob' and knob == 'B':
                base += 6
            noise = rng.normal()
            for shot in range(5):
                stamp = NOW - pd.Timedelta(days=20 - 2 * lot)
                rows.append({'root_lot_id': f'R{lot}', 'fab_lot_id': f'R{lot}.1', 'wafer_id': wafer,
                             '_value': base + rng.normal(0, .5), '_recent': lot >= 8, '_dc_time': stamp, '_time': stamp,
                             '_knob': 'ALL', 'chip_x_pos': shot - 2, 'chip_y_pos': (shot % 2) - 1,
                             '__ml_INLINE_CD': cd, '__ml_INLINE_OX': ox, '__ml_KNOB_ETCH': knob,
                             '__ml_VM_NOISE': noise, '__ml_MASK_REV': 'M1' if lot % 2 else 'M2'})
    return pd.DataFrame(rows)


def entry(effect, seed=1):
    frame = wafer_frame(effect, seed)
    return dict(vehicle='T', item=effect.upper(), category='CAT', step='S1', program='P1', temperature='25', unit='V',
                spatial=frame, points=frame, n=len(frame), lots=10, recent_lots=2, warnings=[], ml_findings=[],
                auto_findings=[], selection_reason='test', aggregation='raw shot', x_label='t', knob_label='', low=None, high=None)


def flagged(e):
    return {(r['column'], s) for r in e['_factors']['rows'] for s in r['signals']}


@pytest.mark.parametrize('effect,expected', [
    ('r2', {('INLINE_CD', 'r2')}),
    ('tail', {('INLINE_OX', 'low_tail')}),
    ('knob', {('KNOB_ETCH', 'level')}),
    ('none', set()),
])
def test_factor_screen_finds_only_the_injected_factor(effect, expected):
    e = entry(effect)
    mf.ml_factor_screen([e], {})
    found = flagged(e)
    assert expected <= found and {c for c, _ in found} <= {c for c, _ in expected}   # 다른 인자는 신호 없음
    if effect == 'tail':   # 한쪽 꼬리만 움직였다 — 상단 꼬리 신호는 없어야 한다
        assert ('INLINE_OX', 'high_tail') not in found
    families = {r['column']: r['family'] for r in e['_factors']['rows']}
    assert families['INLINE_CD'] == 'INLINE' and families['VM_NOISE'] == 'VM' and families['KNOB_ETCH'] == 'KNOB'
    kinds = {r['column']: r['kind'] for r in e['_factors']['rows']}
    assert kinds['INLINE_OX'] == 'numeric' and kinds['KNOB_ETCH'] == 'categorical'


def test_lot_level_factor_is_compared_between_lots_not_wafers():
    e = entry('none')
    mf.ml_factor_screen([e], {})
    mask = next(r for r in e['_factors']['rows'] if r['column'] == 'MASK_REV')
    knob = next(r for r in e['_factors']['rows'] if r['column'] == 'KNOB_ETCH')
    assert mask['unit'].startswith('lot 간') and knob['unit'].startswith('lot 내')


def test_mail_guard_keeps_images_plus_attachments_within_limit():
    content = '<html>' + ''.join(f'<p>{i}</p><img src="data:image/png;base64,AAAA{i}">' for i in range(12)) + '</html>'
    safe, before = main._mail_attachment_guard(content, 1, {'mail_attach_limit': 10})
    assert before == 12 and len(re.findall(r'<img\s', safe)) == 9 and '그림 생략' in safe
    same, _ = main._mail_attachment_guard('<img src="data:image/png;base64,AA">', 1, {'mail_attach_limit': 10})
    assert same.count('<img') == 1


def test_ml_report_pack_splits_mails_by_image_budget(monkeypatch):
    monkeypatch.setattr(main.GLOBAL_CONFIG, 'load_from_yaml', lambda vehicle: None)
    entries = [entry(effect, seed) for seed, effect in enumerate(['r2', 'tail', 'knob', 'none', 'r2', 'none'])]
    for i, e in enumerate(entries):
        e['item'] = f'ITEM{i}'
    settings = dict(service='mlmode', report_now=NOW, chart_dpi=120, mail_image_limit=3)
    mf.ml_factor_screen(entries, settings)
    parts = main._ml_report_pack(entries, settings, 'ML test')
    assert sum(count for _, _, count in parts) == len(entries)
    for body, ppt, count in parts:
        assert len(re.findall(r'<img\s', body)) <= 3
        assert 'ML_TABLE 인자 스크리닝' in body and 'AUTO REPORT · ML MODE' in body
        deck = Presentation(io.BytesIO(ppt))
        # 요약 1장 + 항목마다 (항목 페이지 + 인자 스크리닝 페이지)
        assert len(deck.slides) >= 1 + 2 * count
        titles = [sh.text_frame.text for s in deck.slides for sh in s.shapes if sh.has_text_frame]
        assert any('인자 스크리닝' in t for t in titles)


def test_all_trend_sheets_respect_mail_image_limit():
    from PIL import Image
    tiles = [('CAT', f'ITEM{i}', Image.new('RGB', (500, 250), 'white')) for i in range(40)]
    sheets = main._all_trend_sheets(tiles)
    assert len(sheets) <= main._mail_image_limit() and sum(len(n) for _, n, _ in sheets) == 40
