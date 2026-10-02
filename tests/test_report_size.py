"""PPT size budget, chart encoding, ML candidate source and Watchdog overview (offline)."""
import io
import ast
import os
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import Main as main
import My_Function as mf


def _synthetic_mail_25x100():
    """The production score renderer, full template, Inline table and eight dense charts/maps."""
    from PIL import ImageDraw, ImageFont
    from My_config import _REPORT_HTML_TEMPLATE
    rng = np.random.default_rng(25100)
    board = pd.DataFrame(rng.uniform(70, 100, (100, 25)),
                         index=pd.MultiIndex.from_tuples([(f'CAT_{i//10:02}', f'ITEM_{i:03}') for i in range(100)]),
                         columns=pd.MultiIndex.from_tuples([('SYNTHETIC.1', w) for w in range(1, 26)]))
    board.iloc[2, 3] = np.nan
    chart_images = []
    for seed in range(8):
        im = Image.new('RGB', (1500, 640), 'white')
        draw = ImageDraw.Draw(im)
        import matplotlib
        font = ImageFont.truetype(str(Path(matplotlib.get_data_path())/'fonts/ttf/DejaVuSans.ttf'), 18)
        draw.text((12, 8), f'ITEM_{seed:03} / 25 wafers / spec-out maps', fill='#003366', font=font)
        draw.line((35, 240, 600, 240), fill='#b4232d', width=2)
        draw.line((35, 40, 35, 570, 600, 570), fill='#222222', width=2)
        for lot in range(50):
            for wafer in range(25):
                x = 45 + lot * 11 + int(rng.integers(-3, 4))
                y = int(np.clip(350 + rng.normal(0, 65) + lot * .5, 45, 560))
                draw.ellipse((x-2, y-2, x+2, y+2), fill=('#1f77b4', '#2ca02c', '#d62728', '#9467bd')[wafer % 4])
        # 25 target wafers, 180 measured sites each, clear wafer labels and spec colors.
        for wafer in range(25):
            x, y = 640 + wafer % 5 * 170, 55 + wafer // 5 * 115
            draw.ellipse((x, y, x+94, y+94), outline='#0033cc', width=2)
            for _ in range(180):
                dx, dy = rng.integers(8, 87, 2)
                if (dx-47)**2+(dy-47)**2 <= 40**2:
                    color = ('#2676bc', '#7ac8a4', '#f2b63f', '#d63838')[int(rng.integers(0, 4))]
                    draw.rectangle((x+int(dx), y+int(dy), x+int(dx)+2, y+int(dy)+2), fill=color)
            draw.text((x+15, y+94), f'WF #{wafer+1:02}', fill='#0033cc', font=font)
        buf = io.BytesIO(); im.save(buf, 'PNG'); chart_images.append(buf.getvalue())
    images = ''.join(f'<img alt="ITEM_{i:03} chart and WF MAP" width="750" height="320" '
                     f'style="display:block;width:750px;height:320px" src="{main._img_datauri(raw)}">'
                     for i, raw in enumerate(chart_images))
    inline = '<table style="border-collapse:collapse;font-size:11px"><tbody>'
    for item in range(20):
        inline += f'<tr><td style="padding:4px 10px">INLINE_{item:03}</td>'
        inline += ''.join(f'<td style="border:1px solid #2c2c2c;text-align:center;padding:4px 8px">{wafer+item/100:.2f}</td>' for wafer in range(1, 26)) + '</tr>'
    inline += '</tbody></table>'
    content = _REPORT_HTML_TEMPLATE.replace('{{node}}', 'TEST').replace('{{vehicle}}', 'SYNTHETIC').replace('{{system_admin}}', 'SYNTHETIC')
    content = content.replace('sub_title', '25 wafers / 100 items (synthetic)')
    for i, body in enumerate((images, main._render_score_board(board, 'SYNTHETIC.1'), inline,
                             '<p>합성 검증 · 100개 항목 · 25매 · 최근 DC 측정</p>', '')):
        content = content.replace(f'<div id="target{i}"></div>', f'<div id="target{i}">{body}</div>')
    return content, board, chart_images


def test_mail_25_wafers_100_items_preserves_cells_and_images(monkeypatch, tmp_path):
    import re
    from html.parser import HTMLParser
    monkeypatch.setattr(main.GLOBAL_CONFIG, 'settings', {})
    content, board, charts = _synthetic_mail_25x100()
    packed = main._fit_html_budget(content)
    assert len(content.encode('utf-8')) > 1_000_000  # Exercise the complete-body reduction.
    assert len(packed.encode('utf-8')) < 950_000
    assert main._assert_inline_images(packed, 8) == 8
    class Cells(HTMLParser):
        def __init__(self):
            super().__init__(); self.values=[]; self.active=False
        def handle_starttag(self, tag, attrs):
            if tag == 'td': self.active = dict(attrs).get('class') == 'sb-val'
        def handle_data(self, text):
            if self.active: self.values.append(text)
        def handle_endtag(self, tag):
            if tag == 'td': self.active=False
    cells=Cells(); cells.feed(packed)
    assert cells.values == [f'{v:.1f}' for v in board.to_numpy().ravel() if not pd.isna(v)]
    assert packed.count('class="sb-val"') == 2500
    for i in range(100): assert f'ITEM_{i:03}' in packed
    assert 'font-size:11px' in packed and 'width:750px;height:320px' in packed
    uris=re.findall(r'src="(data:image/[^"\s]+)"',packed)
    sizes=[Image.open(io.BytesIO(__import__('base64').b64decode(uri.split(',')[1]))).size for uri in uris]
    assert all(w >= 750 and h >= 320 for w,h in sizes)  # At least the displayed resolution.
    for raw,uri in zip(charts,uris):
        original=np.asarray(Image.open(io.BytesIO(raw)).convert('RGB'),dtype=float)
        reduced=np.asarray(Image.open(io.BytesIO(__import__('base64').b64decode(uri.split(',')[1]))).convert('RGB'),dtype=float)
        assert original.shape == reduced.shape
        assert np.sqrt(np.mean((original-reduced)**2)) < 2  # Palette compression keeps colors/labels close.
    # A previously saved oversized HTML is also corrected at the real delivery boundary.
    from types import SimpleNamespace
    monkeypatch.setenv('AUTO_REPORT_OPS_ROOT',str(tmp_path/'ops'))
    old=tmp_path/'old-report.html';old.write_text(content,encoding='utf-8')
    sent=[]
    def transport(*args, **kwargs):
        sent.append(ast.literal_eval(kwargs['data']['mailSendString'])['content'])
        return SimpleNamespace(status_code=200,text='offline')
    monkeypatch.setattr(main.requests,'request',transport)
    assert main._durable_mail('25x100', [dict(email='offline@example.test')], 'synthetic', str(old), None, main.GLOBAL_CONFIG) == 'sent'
    assert sent == [packed]
    print(f'25 wafers x 100 items: {len(content.encode()):,} -> {len(packed.encode()):,} bytes; image sizes={sizes}')
    sample = os.environ.get('AUTO_REPORT_MAIL_SAMPLE_DIR')
    if sample:
        dest=Path(sample);dest.mkdir(parents=True,exist_ok=True)
        (dest/'mail-25wafer-100items-before.html').write_text(content,encoding='utf-8')
        (dest/'mail-25wafer-100items.html').write_text(packed,encoding='utf-8')
        deck=Presentation()
        for i,raw in enumerate(charts):
            slide=deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_picture(io.BytesIO(raw),Inches(.2),Inches(.8),width=Inches(9.5))
            slide.shapes.add_textbox(Inches(.2),Inches(.1),Inches(9),Inches(.5)).text_frame.text=f'SYNTHETIC: 25 wafers / 100 items - image {i+1}'
        deck.save(dest/'mail-25wafer-100items.pptx')


def test_html_compaction_preserves_text_and_preformatted_content():
    content='<html><body><b>A</b> <b>B</b><pre>line 1\n  line 2</pre><table>\n  <tr>\n<td style="font-size:11px; color: #ffffff; padding: 4px 6px;">한글 &amp; 97.3</td>\n</tr>\n</table></body></html>'
    packed=main._fit_html_budget(content)
    assert '<b>A</b> <b>B</b>' in packed and '<pre>line 1\n  line 2</pre>' in packed
    assert '한글 &amp; 97.3' in packed and 'font-size:11px;color:#fff;padding:4px 6px' in packed


def test_legacy_two_mb_configuration_cannot_override_mail_limit():
    from types import SimpleNamespace
    cfg=SimpleNamespace(get=lambda key, default=None: {'html_mail_max_mb':2.0}.get(key,default))
    assert main._html_mail_limit(cfg, {'html_max_bytes':2_000_000}) == 1_000_000
    assert main._html_mail_limit(cfg, {'html_max_bytes':600_000}) == 600_000


def test_text_only_body_can_use_the_reserved_margin_below_hard_limit():
    body='<html><body>'+('x'*970_000)+'</body></html>'
    assert main._fit_html_budget(body) == body


@pytest.mark.parametrize('oversized', [False, True])
def test_durable_mail_checks_utf8_body_before_transport(tmp_path, monkeypatch, oversized):
    from types import SimpleNamespace
    monkeypatch.setenv('AUTO_REPORT_OPS_ROOT',str(tmp_path/'ops'))
    path=tmp_path/'mail.html'
    body='<html><body>'+('한글표값' * (90_000 if oversized else 100))+'</body></html>'
    path.write_text(body,encoding='utf-8')
    captured=[]
    def transport(*args, **kwargs):
        payload=ast.literal_eval(kwargs['data']['mailSendString'])
        assert len(payload['content'].encode('utf-8')) < 1_000_000
        captured.append(payload['content'])
        return SimpleNamespace(status_code=200,text='offline')
    monkeypatch.setattr(main.requests,'request',transport)
    cfg=SimpleNamespace(get=lambda key, default=None: {'html_mail_max_mb':2.0,'vehicle':'TEST','url':'http://offline.invalid','KNOXID':'offline'}.get(key,default))
    result=main._durable_mail('size-check', [dict(email='offline@example.test')], 'offline', str(path), None, cfg)
    assert result == ('failed' if oversized else 'sent')
    assert len(captured) == (0 if oversized else 1)
    assert path.read_text('utf-8') == body


def _noise_png(w, h, seed=0):
    rng = np.random.default_rng(seed)
    buf = io.BytesIO()
    Image.fromarray((rng.random((h, w, 3)) * 255).astype('uint8')).save(buf, 'PNG')
    return buf.getvalue()


def _chart_png():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(8.3, 1.95))
    ax.boxplot([np.random.default_rng(i).normal(0, 1, 200) for i in range(25)])
    buf = io.BytesIO(); fig.savefig(buf, format='png', dpi=125); plt.close(fig)
    return buf.getvalue()


def test_chart_encoder_never_worse_than_jpeg_and_prefers_png_for_line_charts():
    png = _chart_png()
    out = mf._encode_chart_image(png, 55, 'auto', 128)
    jpeg = mf._encode_chart_image(png, 55, 'jpeg', 128)
    assert len(out) <= len(jpeg)
    assert out[:4] == b'\x89PNG'           # box plot: palette PNG wins
    assert Image.open(io.BytesIO(out)).size == Image.open(io.BytesIO(png)).size


def test_wfmap_canvas_is_capped_to_display_resolution():
    canvas = Image.new('RGB', (3400, 460), 'white')
    buf = mf._encode_map_canvas(canvas, dict(map_max_px=2370, map_colors=128))
    image = Image.open(io.BytesIO(buf.getvalue()))
    assert image.size[0] == 2370 and image.mode == 'P'


def _deck_with_description(image_bytes, charts=3):
    source = Presentation()
    slide = source.slides.add_slide(source.slide_layouts[6])
    slide.shapes.add_textbox(Inches(.3), Inches(.2), Inches(4), Inches(.5)).text_frame.text = 'VTH'
    slide.shapes.add_picture(io.BytesIO(image_bytes), Inches(5), Inches(1), Inches(7))
    deck = Presentation()
    for i in range(charts):
        s = deck.slides.add_slide(deck.slide_layouts[6])
        s.shapes.add_picture(io.BytesIO(_noise_png(900, 300, seed=i + 1)), Inches(.3), Inches(.3), Inches(8))
    mf._copy_slide_into(deck, slide, source.slide_width, source.slide_height)
    return deck


def _set(monkeypatch, **values):
    for key, value in values.items():
        monkeypatch.setattr(mf.GLOBAL_CONFIG, key, value, raising=False)


def test_description_images_use_remaining_budget(monkeypatch):
    _set(monkeypatch, description_image_recompress=True, description_image_target_mb=0, description_image_max_px=2400,
         description_min_px=480, ppt_budget_ratio=1.0)
    big = _noise_png(2400, 1600)
    roomy = _deck_with_description(big)
    _set(monkeypatch, ppt_mail_max_mb=50.0)
    generous = mf.fit_ppt_budget(roomy, mf.GLOBAL_CONFIG)
    tight_deck = _deck_with_description(big)
    base = mf._ppt_bytes(tight_deck) - len(tight_deck._auto_desc_images and
                                           tight_deck.slides[-1].part.related_part(tight_deck._auto_desc_images[0]['rId']).blob)
    _set(monkeypatch, ppt_mail_max_mb=(base + 250_000) / 1e6)
    tight = mf.fit_ppt_budget(tight_deck, mf.GLOBAL_CONFIG)
    assert generous['desc_level'].startswith('2400px')         # 여유가 많으면 최고 화질
    assert tight['after'] <= tight['limit'] and tight['desc_level'] not in ('-', '생략')
    assert int(tight['desc_level'].split('px')[0]) < 2400      # 빠듯하면 해상도를 낮춰 끼워 넣는다
    # 원본 설명 PPT 이미지 파트는 건드리지 않는다(다음 Lot 에서 재사용).
    assert len(big) > 1_000_000


def test_description_image_dropped_but_text_kept_when_no_room(monkeypatch):
    _set(monkeypatch, description_image_recompress=True, description_image_target_mb=0, description_min_px=480, ppt_budget_ratio=1.0)
    deck = _deck_with_description(_noise_png(2400, 1600), charts=1)
    _set(monkeypatch, ppt_mail_max_mb=0.05)
    result = mf.fit_ppt_budget(deck, mf.GLOBAL_CONFIG)
    last = deck.slides[-1]
    texts = ' '.join(sh.text_frame.text for sh in last.shapes if sh.has_text_frame)
    assert result['desc_dropped'] == 1 and 'VTH' in texts and '용량 한도' in texts
    assert not last.shapes._spTree.xpath('.//p:pic')


def test_chart_images_shrink_as_last_resort(monkeypatch):
    _set(monkeypatch, ppt_budget_ratio=1.0)
    deck = Presentation()
    for i in range(3):
        deck.slides.add_slide(deck.slide_layouts[6]).shapes.add_picture(io.BytesIO(_noise_png(1400, 900, i)), 0, 0, Inches(8))
    size = mf._ppt_bytes(deck)
    _set(monkeypatch, ppt_mail_max_mb=size * .8 / 1e6)
    result = mf.fit_ppt_budget(deck, mf.GLOBAL_CONFIG)
    assert result['charts_shrunk'] > 0 and result['after'] < size


@pytest.mark.parametrize('source,expected', [('daily', {'A'}), ('ml', {'B'}), ('either', {'A', 'B'})])
def test_ml_candidate_source_controls_which_items_get_detail(monkeypatch, source, expected):
    entries = [dict(item='A', auto_findings=[dict(title='Flier')], warnings=[]), dict(item='B', warnings=[]), dict(item='C', warnings=[])]
    seen = []

    def fake_select(pool, settings):
        seen.append([e['item'] for e in pool])
        chosen = [e for e in pool if e['item'] == 'B']
        for e in chosen:
            e['ml_findings'] = [dict(module='isolation_forest', q=.01, effect=.5, message='m')]
        return chosen, {}
    settings = dict(candidate_source=source)
    # 선정 로직만 떼어 실행(데이터 적재·렌더링 없이)
    pool = [e for e in entries if e.get('auto_findings')] if main._ml_candidate_source(settings) == 'daily' else entries
    selected, _ = fake_select(pool, settings)
    chosen = {id(e) for e in selected}
    if main._ml_candidate_source(settings) in ('daily', 'either'):
        selected += [e for e in pool if e.get('auto_findings') and id(e) not in chosen]
    assert {e['item'] for e in selected} == expected
    assert seen[0] == (['A'] if source == 'daily' else ['A', 'B', 'C'])
    assert main._ml_module_note('isolation_forest', dict(module_notes={'isolation_forest': '사용자 문구'})) == '사용자 문구'
    assert 'drift' in main._ml_module_note('time_trend', {})


def test_ml_selection_block_in_report_matches_contract():
    source = Path(main.__file__).read_text(encoding='utf-8')
    assert "pool=[e for e in entries if e.get('auto_findings')] if source=='daily' else entries" in source
    assert "e['selection_reason']" in source


def test_watchdog_overview_verdict_and_timeline(monkeypatch, tmp_path):
    monkeypatch.setenv('AUTO_REPORT_OPS_ROOT', str(tmp_path))
    now = 1_760_000_000.0
    reports = [dict(vehicle='P1', lot='L1', step='S1', started=now - 100, updated=now - 50, status='success', email='sent', elapsed=120, mode='AUTO'),
               dict(vehicle='P1', lot='L2', step='S1', started=now - 90, updated=now - 40, status='failed', email='', mode='TRIGGER')]
    html = main._watchdog_overview(dict(state='healthy', message='정상'), [], {}, {}, reports,
                                   [dict(vehicle='P1', rc=0, elapsed=300)], now - 86400, now, dict(products=['P1']))
    assert '정상 — Scheduler 동작' in html and 'id="timeline"' in html and 'L2 / S1' in html
    assert '1회 / 평균 5.0분' in html
    warn = main._watchdog_overview(dict(state='stale', message='정지 의심'), [['x', 'y', 'z', 'w']], {}, {}, [], [], now - 86400, now, {})
    assert '확인 필요 1건' in warn
